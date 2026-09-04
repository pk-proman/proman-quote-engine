"""
PROMAN Offer Letter Generator
Produces a Word (.docx) offer letter with:
  Cover page + Annexure 1 (Summary) + Annexure 3 (Tech Specs)
  + Annexure 4 (Price Schedule narrative) + Annexure 5 (T&C) + Annexure 6 (Exclusions)
Requires: python-docx
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys, os, copy
from datetime import date as dt_date

sys.path.insert(0, os.path.dirname(__file__))
from engine import Quotation, PlantSpec, build_quotation, quote_standalone_equipment
import rate_card as rc


# ---------------------------------------------------------------------------
# COLOURS
# ---------------------------------------------------------------------------
PROMAN_BLUE  = RGBColor(0x2E, 0x75, 0xB6)
DARK_NAVY    = RGBColor(0x1F, 0x38, 0x64)
LIGHT_BLUE   = RGBColor(0xBD, 0xD7, 0xEE)
YELLOW       = RGBColor(0xF2, 0xCC, 0x0C)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
BLACK        = RGBColor(0x00, 0x00, 0x00)


# ---------------------------------------------------------------------------
# LOW-LEVEL HELPERS
# ---------------------------------------------------------------------------

def _set_cell_bg(cell, hex_color: str):
    """Set table cell background colour (OOXML shading)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _para(doc, text, bold=False, size=11, color=None, align=WD_ALIGN_PARAGRAPH.LEFT,
          space_before=0, space_after=4, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold  = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def _heading(doc, text, level=1):
    sizes = {1: 14, 2: 12, 3: 11}
    p = _para(doc, text, bold=True, size=sizes.get(level, 11),
              color=DARK_NAVY, space_before=8, space_after=4)
    return p


def _page_break(doc):
    doc.add_page_break()


def _section_rule(doc):
    """Thin horizontal line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E75B6")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _table_row(table, values, bold=False, bg_hex=None, col_widths=None):
    """Add a row to a table with optional bold and background colour."""
    row = table.add_row()
    for ci, (cell, val) in enumerate(zip(row.cells, values)):
        cell.text = str(val) if val is not None else ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for para in cell.paragraphs:
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after  = Pt(1)
            for run in para.runs:
                run.bold = bold
                run.font.size = Pt(9)
                if bg_hex in ("1F3864", "2E75B6"):
                    run.font.color.rgb = WHITE
        if bg_hex:
            _set_cell_bg(cell, bg_hex)
    if col_widths:
        for ci, w in enumerate(col_widths):
            row.cells[ci].width = Cm(w)
    return row


# ---------------------------------------------------------------------------
# DOCUMENT SECTIONS
# ---------------------------------------------------------------------------

def _cover_page(doc: Document, q: Quotation):
    spec = q.spec
    today = spec.date or dt_date.today().strftime("%d %B %Y")

    # Company header
    p = _para(doc, "PROMAN INFRASTRUCTURE SERVICES PVT. LTD.",
              bold=True, size=16, color=DARK_NAVY,
              align=WD_ALIGN_PARAGRAPH.CENTER, space_before=30, space_after=2)
    _para(doc, "No. 7/1, 17th Cross, Sadashivanagar, Bengaluru – 560 080",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2)
    _para(doc, "Tel: +91-80-4112 2277  |  Email: info@promaninfra.com  |  www.promaninfra.com",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=10)
    _section_rule(doc)

    # Ref & date block
    ref_table = doc.add_table(rows=2, cols=2)
    ref_table.style = "Table Grid"
    ref_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ref_cells = [[ref_table.rows[0].cells[0], ref_table.rows[0].cells[1]],
                 [ref_table.rows[1].cells[0], ref_table.rows[1].cells[1]]]
    ref_cells[0][0].text = f"Ref. No.: {spec.ref_no}"
    ref_cells[0][1].text = f"Date: {today}"
    ref_cells[1][0].text = f"To: M/s. {spec.client_name}"
    ref_cells[1][1].text = ""
    for r in ref_table.rows:
        for c in r.cells:
            for para in c.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()

    # Subject
    subj = f"Offer for Supply of {spec.tph} TPH, {spec.stages}-Stage, "
    subj += "Portable" if spec.mobile else "Stationary"
    subj += " Aggregate Crushing & Screening Plant"
    _para(doc, f"Subject: {subj}", bold=True, size=11, space_before=10)

    # Salutation
    _para(doc, "Dear Sir/Madam,", size=11, space_before=6, space_after=4)

    body = (
        f"With reference to your enquiry and our subsequent discussions, we are pleased "
        f"to submit our techno-commercial offer for the supply of the above-mentioned "
        f"plant and equipment.\n\n"
        f"We trust that our offer meets your requirements. We request you to go through the "
        f"techno-commercial details and look forward to your valued order.\n\n"
        f"This offer is subject to terms & conditions mentioned in Annexure 5."
    )
    _para(doc, body, size=11, space_before=0, space_after=10)
    _para(doc, "Thanking you,", size=11)
    _para(doc, "For PROMAN Infrastructure Services Pvt. Ltd.", size=11, bold=True,
          space_after=30)
    _para(doc, "Authorised Signatory", size=10)


def _annexure_1_summary(doc: Document, q: Quotation):
    _page_break(doc)
    _para(doc, "ANNEXURE 1 — QUOTATION SUMMARY", bold=True, size=13,
          color=DARK_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _section_rule(doc)

    spec = q.spec
    rows = [
        ("Client",                    f"M/s. {spec.client_name}"),
        ("Plant Capacity",            f"{spec.tph} TPH"),
        ("Plant Type",                f"{spec.stages}-Stage, {'Portable/Skid' if spec.mobile else 'Stationary'}"),
        ("Feed Material",             spec.material),
        ("Max. Feed Size",            f"{spec.feed_size_mm} mm"),
        ("Output Products",           ", ".join(spec.products) if spec.products else "As per client requirement"),
        ("Sub A — Equipment (List)",  f"₹ {q.sub_a_list():.2f} Lakhs"),
        ("Sub B — Conveyors (List)",  f"₹ {q.sub_b_total_list():.2f} Lakhs"),
        ("Sub C — Structural (List)", f"₹ {q.sub_c_structural_list:.2f} Lakhs"),
        ("Sub D — Electrical (List)", f"₹ {q.sub_d_list:.2f} Lakhs"),
        ("Total A+B+C+D (List)",      f"₹ {q.total_list():.2f} Lakhs"),
        ("Sub E — Erection",          f"₹ {q.erection:.2f} Lakhs"),
        ("GRAND TOTAL (List)",        f"₹ {q.grand_total_list():.2f} Lakhs"),
        ("GRAND TOTAL (Best)",        f"₹ {q.grand_total_best():.2f} Lakhs"),
        ("Total Connected HP",        f"{q.total_hp()} HP"),
        ("Delivery",                  "12–14 weeks from receipt of technically & commercially clear PO + advance"),
        ("Validity",                  "30 days from date of offer"),
        ("Payment",                   "40% advance with PO; 60% + 100% GST against Proforma Invoice before dispatch"),
    ]

    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _table_row(tbl, ["Parameter", "Details"], bold=True, bg_hex="2E75B6")
    for i, (k, v) in enumerate(rows):
        bg = "BDD7EE" if i % 2 == 0 else None
        bold = "GRAND" in k or "TOTAL" in k
        _table_row(tbl, [k, v], bold=bold, bg_hex=("1F3864" if "GRAND TOTAL" in k else bg))

    if q.warnings:
        doc.add_paragraph()
        _para(doc, "⚠️  Items Requiring Confirmation", bold=True, size=10, color=RGBColor(0xC0,0,0))
        for w in q.warnings:
            _para(doc, f"  •  {w}", size=9, color=RGBColor(0xC0,0,0), space_after=2)


def _annexure_3_tech_spec(doc: Document, q: Quotation):
    _page_break(doc)
    _para(doc, "ANNEXURE 3 — TECHNICAL SPECIFICATIONS", bold=True, size=13,
          color=DARK_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _section_rule(doc)

    spec = q.spec
    _para(doc, "3.1  Plant Parameters", bold=True, size=11, color=PROMAN_BLUE,
          space_before=6, space_after=3)

    param_rows = [
        ("Capacity",          f"{spec.tph} TPH"),
        ("No. of Stages",     str(spec.stages)),
        ("Feed Material",     spec.material),
        ("Max. Feed Size",    f"{spec.feed_size_mm} mm"),
        ("Products",          ", ".join(spec.products) if spec.products else "As per requirement"),
        ("Plant Type",        "Portable / Skid-mounted" if spec.mobile else "Stationary"),
        ("Installed Power",   f"Approx. {q.total_hp()} HP Connected Load"),
    ]

    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    _table_row(tbl, ["Parameter", "Specification"], bold=True, bg_hex="2E75B6")
    for i, (k, v) in enumerate(param_rows):
        _table_row(tbl, [k, v], bg_hex=("BDD7EE" if i % 2 == 0 else None))

    doc.add_paragraph()
    _para(doc, "3.2  Equipment Schedule", bold=True, size=11, color=PROMAN_BLUE,
          space_before=6, space_after=3)

    eq_tbl = doc.add_table(rows=1, cols=4)
    eq_tbl.style = "Table Grid"
    _table_row(eq_tbl, ["Sl.", "Equipment", "Qty", "Drive HP"], bold=True, bg_hex="2E75B6")
    for i, item in enumerate(q.sub_a_items):
        _table_row(eq_tbl, [item.sl, item.description, item.qty, item.hp],
                   bg_hex=("BDD7EE" if i % 2 == 0 else None))

    doc.add_paragraph()
    _para(doc, "3.3  Belt Conveyor Schedule", bold=True, size=11, color=PROMAN_BLUE,
          space_before=6, space_after=3)
    conv_tbl = doc.add_table(rows=1, cols=5)
    conv_tbl.style = "Table Grid"
    _table_row(conv_tbl, ["Tag", "Width (mm)", "Length (m)", "Drive HP", "Purpose"],
               bold=True, bg_hex="2E75B6")
    for i, c in enumerate(q.sub_b_conveyors):
        _table_row(conv_tbl,
                   [c.tag, c.width_mm, f"{c.length_m:.1f}", c.hp, c.purpose.replace("_", " ")],
                   bg_hex=("BDD7EE" if i % 2 == 0 else None))


def _annexure_4_price_schedule(doc: Document, q: Quotation):
    """Narrative price schedule in Word (companion to the Excel Annexure 4)."""
    _page_break(doc)
    _para(doc, "ANNEXURE 4 — PRICE SCHEDULE", bold=True, size=13,
          color=DARK_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _para(doc, "(Detailed Excel price schedule provided as a separate file)",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x59,0x59,0x59),
          space_after=6)
    _section_rule(doc)

    # Summary table only
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    _table_row(tbl, ["Sub", "Description", "Amount (₹ Lakhs — List)"],
               bold=True, bg_hex="2E75B6")
    data = [
        ("A", "Main Mechanical Equipment",       f"{q.sub_a_list():.2f}"),
        ("B", "Belt Conveyor System",             f"{q.sub_b_total_list():.2f}"),
        ("C", "Structural Steel Fabricated Items",f"{q.sub_c_structural_list:.2f}"),
        ("D", "Electrical Scope of Work",         f"{q.sub_d_list:.2f}"),
        ("",  "Total  (A+B+C+D)",                f"{q.total_list():.2f}"),
        ("E", "Erection & Commissioning",         f"{q.erection:.2f}"),
        ("",  "GRAND TOTAL",                      f"{q.grand_total_list():.2f}"),
    ]
    for i, (s, d, v) in enumerate(data):
        bold = "GRAND" in d or "Total" in d
        bg = "1F3864" if "GRAND" in d else ("BDD7EE" if i % 2 == 0 else None)
        _table_row(tbl, [s, d, v], bold=bold, bg_hex=bg)

    doc.add_paragraph()
    _para(doc, "Note: GST as applicable at time of dispatch is extra. Prices are Ex-Works, Bengaluru.",
          size=9, color=RGBColor(0x59,0x59,0x59))


def _annexure_5_terms(doc: Document):
    _page_break(doc)
    _para(doc, "ANNEXURE 5 — TERMS & CONDITIONS", bold=True, size=13,
          color=DARK_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _section_rule(doc)

    terms = [
        ("Payment",        "40% non-refundable advance with Purchase Order. Balance 60% + 100% GST against "
                           "Proforma Invoice prior to dispatch of each consignment."),
        ("Price Basis",    "Ex-Works, Bengaluru. Prices are in Indian Rupees and are firm for the validity period."),
        ("GST",            "GST as applicable at the time of dispatch will be charged extra."),
        ("Freight",        "Freight, transit insurance and unloading at site are in client's scope and extra."),
        ("Delivery",       "12 to 14 weeks from receipt of technically and commercially clear Purchase Order "
                           "along with advance. Subject to Force Majeure."),
        ("Validity",       "This offer is valid for 30 days from date of issue."),
        ("Warranty",       "12 months from the date of commissioning or 18 months from the date of dispatch, "
                           "whichever is earlier. Warranty covers manufacturing defects only."),
        ("Erection",       "PROMAN engineers will supervise erection. Client to arrange crane, "
                           "skilled and unskilled labour, consumables, and access to site."),
        ("Bank Guarantee", "Performance Bank Guarantee is not applicable for this order."),
        ("Jurisdiction",   "Any dispute shall be subject to jurisdiction of courts in Bengaluru only."),
    ]
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    _table_row(tbl, ["Clause", "Detail"], bold=True, bg_hex="2E75B6")
    for i, (k, v) in enumerate(terms):
        _table_row(tbl, [k, v], bg_hex=("BDD7EE" if i % 2 == 0 else None))


def _annexure_6_exclusions(doc: Document):
    _page_break(doc)
    _para(doc, "ANNEXURE 6 — SCOPE EXCLUSIONS", bold=True, size=13,
          color=DARK_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _section_rule(doc)
    _para(doc, "The following items are NOT included in this offer:", bold=True, size=10,
          space_before=6, space_after=4)

    exclusions = [
        "Civil work including foundations, flooring, retaining walls, site levelling, and drainage.",
        "Power supply connection up to MCC panel (HT/LT line, transformer, DG set).",
        "Earthing and lightning protection.",
        "Compressed air / water supply lines.",
        "Dust suppression system and water sprinklers.",
        "Weighbridge and weigh-in-motion systems.",
        "Lighting inside the plant area.",
        "Fencing, boundary wall, and gates.",
        "Safety signage and PPE for site workers.",
        "Loader / excavator / dumper for feeding the plant.",
        "Any statutory approvals, royalties, or mine-related permits.",
        "Consumables: wear parts, liners, hammer, jaw plates — for day-to-day operations.",
        "Operator training beyond standard commissioning handover.",
        "Any items not explicitly listed in Annexure 3 and 4.",
    ]
    for ex in exclusions:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(ex)
        run.font.size = Pt(10)


# ---------------------------------------------------------------------------
# STANDALONE QUOTE (for single equipment — Chiman Bhai style)
# ---------------------------------------------------------------------------

def build_standalone_letter(category: str, model: str,
                             client_name: str, ref_no: str,
                             out_path: str) -> str:
    """
    Build a short offer letter for a single equipment.
    e.g. VSI-300 standalone for Chiman Bhai Sapariya.
    """
    data = quote_standalone_equipment(category, model)
    today = dt_date.today().strftime("%d %B %Y")

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Header
    _para(doc, "PROMAN INFRASTRUCTURE SERVICES PVT. LTD.", bold=True, size=15,
          color=DARK_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, "No. 7/1, 17th Cross, Sadashivanagar, Bengaluru – 560 080",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, "Tel: +91-80-4112 2277  |  Email: info@promaninfra.com",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _section_rule(doc)

    # Meta
    _para(doc, f"Ref. No.: {ref_no}    |    Date: {today}", size=10, space_before=8)
    _para(doc, f"To: M/s. {client_name}", size=10, bold=True, space_after=8)
    _para(doc, f"Subject: Offer for {category} {model} — Standalone Equipment Supply",
          bold=True, size=11, space_before=6, space_after=6)
    _para(doc, "Dear Sir/Madam,", size=11, space_after=4)
    _para(doc, (
        f"With reference to your enquiry, we are pleased to submit our offer for supply "
        f"of the REMCO VSI-{model} (if VSI) / {category}-{model} equipment as detailed below."
    ), size=11, space_after=10)

    # Price table
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    _table_row(tbl, ["Description", "List Price (₹ Lakhs)", "Best Price (₹ Lakhs)"],
               bold=True, bg_hex="2E75B6")

    cat_disp = category.upper()
    model_disp = model

    _table_row(tbl, [
        f"REMCO {cat_disp}-{model_disp} with IP Protection" if cat_disp == "VSI"
        else f"{cat_disp}-{model_disp}",
        f"{data['list']:.2f}" if data['list'] else "As per engineering",
        f"{data['best']:.2f}" if data['best'] else "As per engineering",
    ], bg_hex="BDD7EE")

    if data.get("starter_soft"):
        s = data["starter_soft"]
        _table_row(tbl, [
            f"Soft Starter for {data.get('hp', '')} HP Motor",
            f"{s['list']:.2f}" if s.get('list') else "—",
            f"{s['best']:.2f}" if s.get('best') else "—",
        ])
        _table_row(tbl, [
            "Total (Equipment + Soft Starter)",
            f"{data['list'] + s['list']:.2f}" if data['list'] and s.get('list') else "—",
            f"{data['best'] + s['best']:.2f}" if data['best'] and s.get('best') else "—",
        ], bold=True, bg_hex="1F3864")

    if data.get("flags"):
        doc.add_paragraph()
        for f in data["flags"]:
            _para(doc, f"⚠️  {f}", size=9, color=RGBColor(0xC0,0,0))

    # Key terms
    doc.add_paragraph()
    _para(doc, "Key Terms:", bold=True, size=10, space_before=8)
    terms = [
        "Prices are Ex-Works, Bengaluru. GST extra.",
        "Freight & insurance at actuals — client's scope.",
        "Payment: 50% advance with PO; 50% against Proforma Invoice before dispatch.",
        "Delivery: 10–12 weeks from receipt of clear PO + advance.",
        "Warranty: 12 months from commissioning / 18 months from dispatch.",
        "Validity: 30 days from date of this offer.",
    ]
    for t in terms:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(t).font.size = Pt(10)

    doc.add_paragraph()
    _para(doc, "We trust this meets your requirement. Please feel free to reach us for any clarifications.",
          size=11, space_before=6)
    _para(doc, "\nThanking you,\nFor PROMAN Infrastructure Services Pvt. Ltd.", size=11, bold=True,
          space_after=30)
    _para(doc, "Authorised Signatory", size=10)

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# FULL PLANT OFFER LETTER
# ---------------------------------------------------------------------------

def build_offer_letter(q: Quotation, out_path: str) -> str:
    """
    Build the full Word offer letter for a plant quotation.
    Writes: Cover + Annexure 1 + 3 + 4 + 5 + 6
    """
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _cover_page(doc, q)
    _annexure_1_summary(doc, q)
    _annexure_3_tech_spec(doc, q)
    _annexure_4_price_schedule(doc, q)
    _annexure_5_terms(doc)
    _annexure_6_exclusions(doc)

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# CLI — run both Chiman Bhai standalone + 2-stage test
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    out_dir = os.path.dirname(__file__)

    # 1. Chiman Bhai Sapariya — Standalone VSI-300
    p1 = build_standalone_letter(
        category="VSI", model="300",
        client_name="Chiman Bhai Sapariya / SR Petrol Pump",
        ref_no="PR/08-26/N-001",
        out_path=os.path.join(out_dir, "Offer_ChimanBhai_VSI300.docx")
    )
    print(f"✅  Standalone letter: {p1}")

    # 2. Full 2-stage 250TPH plant
    spec = PlantSpec(
        client_name="Test Client — 2 Stage",
        tph=250, stages=2,
        products=["GSB", "-20+10mm", "-10+6mm", "-6+4.75mm", "0-4.75mm"],
        ref_no="PR/08-26/N-TEST",
    )
    q = build_quotation(spec)
    p2 = build_offer_letter(q, os.path.join(out_dir, "Offer_Test_2Stage_250TPH.docx"))
    print(f"✅  Full plant letter: {p2}")
